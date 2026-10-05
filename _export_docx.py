# -*- coding: utf-8 -*-
"""把 topics.json 中已发布的文章导出为可通读的 Word 合集。"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = r"F:\guangergaozhi\每日定时触发"
DEST = os.path.join(BASE, "已发布24篇合集.docx")
SITE = "https://shakhama.github.io/mei-ri-chufa"

d = json.load(open(os.path.join(BASE, 'topics.json'), encoding='utf-8'))
pub = [x for x in d['topics'] if x.get('publish_date')]
pub.sort(key=lambda x: x['publish_date'])


def body_wc(x):
    n = 0
    for s in x.get('sections', []):
        n += len(s.get('body', '') or '')
        for b in s.get('bullets', []) or []:
            n += len(b)
    return n


def full_wc(x):
    n = len(x.get('summary', '') or '')
    for s in x.get('sections', []):
        n += len(s.get('h2', '') or '') + len(s.get('body', '') or '')
        for b in s.get('bullets', []) or []:
            n += len(b)
    for f in x.get('faq', []) or []:
        n += len(f.get('q', '') or '') + len(f.get('a', '') or '')
    return n


bw = [body_wc(x) for x in pub]
fw = [full_wc(x) for x in pub]

doc = Document()

# 页面
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)

# 默认字体
st = doc.styles['Normal']
st.font.name = '微软雅黑'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(4)


def setfont(run, size=None, bold=None, color=None, italic=None):
    f = run.font
    f.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = RGBColor(*color)
    return run


def para(text='', size=10.5, bold=None, color=None, italic=None,
         align=None, before=0, after=4, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.left_indent = Cm(indent)
    if text:
        setfont(p.add_run(text), size, bold, color, italic)
    return p


# ---------- 封面 ----------
para('力哥设计实验室', size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=2)
para('已发布文章合集 · 全 24 篇', size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
para(f"发布区间 {pub[0]['publish_date']} → {pub[-1]['publish_date']}　|　导出于 2026-10-05",
     size=10, color=(0x77, 0x77, 0x77), align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
para('用于整体通读复核：每篇含发布日期、关键词、正文字数、线上地址、全文与 FAQ',
     size=9.5, color=(0x88, 0x88, 0x88), align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

# ---------- 一、内容总览 ----------
para('一、内容总览', size=14, bold=True, before=10, after=6)

tbl = doc.add_table(rows=1, cols=6)
tbl.style = 'Table Grid'
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ['#', '发布日期', '标题', '正文', '全口径', '关键词']
widths = [Cm(0.9), Cm(2.0), Cm(7.2), Cm(1.3), Cm(1.5), Cm(4.0)]
for i, h in enumerate(hdr):
    c = tbl.rows[0].cells[i]
    c.text = ''
    setfont(c.paragraphs[0].add_run(h), size=9, bold=True)
    c.width = widths[i]

for i, x in enumerate(pub, 1):
    row = tbl.add_row()
    kw = '、'.join(x.get('keywords', [])[:3])
    vals = [str(i), x['publish_date'], x['title'], str(bw[i-1]), str(fw[i-1]), kw]
    for j, v in enumerate(vals):
        c = row.cells[j]
        c.text = ''
        setfont(c.paragraphs[0].add_run(v), size=8.5)
        c.width = widths[j]

# ---------- 二、厚度体检 ----------
para('二、厚度体检', size=14, bold=True, before=14, after=6)

avg_b = round(sum(bw) / len(bw), 1)
avg_f = round(sum(fw) / len(fw), 1)
n300 = sum(1 for a in bw if a < 300)
n600 = sum(1 for a in fw if a < 600)

para(f"已发 {len(pub)} 篇，正文（body+bullets）平均 {avg_b} 字，最短 {min(bw)}，最长 {max(bw)}", size=10, after=2)
para(f"全口径（含摘要/小标题/FAQ）平均 {avg_f} 字，最短 {min(fw)}，最长 {max(fw)}", size=10, after=2)
p = para('', size=10, after=2)
setfont(p.add_run(f"正文不足 300 字：{n300} / {len(bw)} 篇"), bold=True, color=(0xC0, 0x39, 0x2B))
p = para('', size=10, after=2)
setfont(p.add_run(f"全口径不足 600 字：{n600} / {len(fw)} 篇"), bold=True, color=(0xC0, 0x39, 0x2B))
p = para('', size=10, after=2)
setfont(p.add_run("参考目标 1200~1800 字/篇，当前均值差 4~6 倍 —— 这是搜索流量路径上的主要卡点。"),
        bold=True, color=(0xC0, 0x39, 0x2B))

# ---------- 三、目录 ----------
doc.add_page_break()
para('三、目录', size=14, bold=True, after=8)
para('正文按发布顺序排列，每篇独立起页。', size=9.5, color=(0x88, 0x88, 0x88), after=8)
for i, x in enumerate(pub, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.4)
    setfont(p.add_run(f"{i:>2}. "), size=10, bold=True)
    setfont(p.add_run(x['title']), size=10)
    setfont(p.add_run(f"　（{x['publish_date']}，{bw[i-1]} 字）"), size=9, color=(0x88, 0x88, 0x88))

# ---------- 正文 ----------
for i, x in enumerate(pub, 1):
    doc.add_page_break()
    para(f"{i}. {x['title']}", size=16, bold=True, before=0, after=6)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    setfont(p.add_run(f"发布日期 {x['publish_date']}　|　slug {x['slug']}　|　"
                      f"正文 {bw[i-1]} 字　|　全口径 {fw[i-1]} 字"), size=9, color=(0x66, 0x66, 0x66))

    kw = '、'.join(x.get('keywords', []))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    setfont(p.add_run('关键词：'), size=9, bold=True, color=(0x66, 0x66, 0x66))
    setfont(p.add_run(kw), size=9, color=(0x66, 0x66, 0x66))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    setfont(p.add_run('线上：'), size=9, bold=True, color=(0x1A, 0x6F, 0xB0))
    setfont(p.add_run(f"{SITE}/{x['slug']}.html"), size=9, color=(0x1A, 0x6F, 0xB0))

    if x.get('summary'):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(8)
        setfont(p.add_run(x['summary']), size=10, italic=True, color=(0x44, 0x44, 0x44))

    for s in x.get('sections', []):
        if s.get('h2'):
            para(s['h2'], size=12.5, bold=True, before=10, after=4)
        if s.get('body'):
            para(s['body'], size=10.5, after=4)
        for b in s.get('bullets', []) or []:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.8)
            setfont(p.add_run(b), size=10.5)

    faq = x.get('faq', []) or []
    if faq:
        para('常见问题', size=12.5, bold=True, before=10, after=4)
        for f in faq:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            setfont(p.add_run('问：'), size=10.5, bold=True)
            setfont(p.add_run(f.get('q', '')), size=10.5, bold=True)
            para(f.get('a', ''), size=10.5, after=6, indent=0.6)

doc.save(DEST)
print("SAVED", DEST, os.path.getsize(DEST), "bytes")
print("pub", len(pub), "avg_body", avg_b, "avg_full", avg_f, "<300:", n300, "<600:", n600)
